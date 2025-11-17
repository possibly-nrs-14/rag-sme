# part_H_tools.py
import os
import logging
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas
from docx import Document
from pptx import Presentation
import textwrap
from langchain_core.tools import BaseTool
from typing import Type, Dict, Any, Optional
from pydantic.v1 import BaseModel, Field
import time
logger = logging.getLogger(__name__)

def export_qa_pdf(question, answer, sources, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    c = canvas.Canvas(out_path, pagesize=LETTER)
    w, h = LETTER
    y = h - 72
    c.setFont("Times-Roman", 14)
    c.drawString(72, y, "Question:")
    y -= 24
    c.setFont("Times-Roman", 12)
    for line in textwrap.wrap(question, width=95):
        c.drawString(72, y, line)
        y -= 18
        if y < 120:
            c.showPage()
            y = h - 72
    y -= 12
    c.setFont("Times-Roman", 14)
    c.drawString(72, y, "Answer:")  
    y -= 24
    c.setFont("Times-Roman", 12)
    for line in textwrap.wrap(answer, width=95):
        c.drawString(72, y, line)
        y -= 18
        if y < 120:
            c.showPage()
            y = h - 72
    y -= 12
    c.setFont("Times-Roman", 14)
    c.drawString(72, y, "Sources:")  
    y -= 24
    c.setFont("Times-Roman", 12)
    for s in sources:
        c.drawString(72, y, f"{s.get('book')} | chunk={s.get('chunk_id')} | g={s.get('granularity')} | pos={s.get('position')}")
        y -= 18
        if y < 120:
            c.showPage()
            y = h - 72
    c.save()

def export_quiz_pdf(topic, items, sources, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    c = canvas.Canvas(out_path, pagesize=LETTER)
    w, h = LETTER
    y = h - 72
    
    # Add topic at the top if provided
    if topic:
        c.setFont("Times-Bold", 16)
        c.drawString(72, y, f"Quiz on {topic}")
        y -= 36
    
    num = 1
    for it in items:
        q = it.get("question") or ""
        opts = it.get("options") or {}
        c.setFont("Times-Bold", 12)
        
        # Wrap question text
        wrapped_q = textwrap.wrap(f"{num}. {q}", width=95)
        for wrapped in wrapped_q:
            c.drawString(72, y, wrapped)
            y -= 18
            if y < 120:
                c.showPage()
                y = h - 72
        
        c.setFont("Times-Roman", 11)
        for k in ["A","B","C","D"]:
            v = opts.get(k, "")
            wrapped_opts = textwrap.wrap(f"{k}. {v}", width=90)
            for wrapped in wrapped_opts:
                c.drawString(90, y, wrapped)
                y -= 16
                if y < 120:
                    c.showPage()
                    y = h - 72
        y -= 6
        num += 1
    
    c.showPage()
    c.setFont("Times-Bold", 12)
    c.drawString(72, h-72, "Answer Key")
    y = h - 96
    c.setFont("Times-Roman", 11)
    for i, it in enumerate(items, 1):
        c.drawString(72, y, f"{i}. {it.get('correct')}")
        y -= 16
        if y < 72:
            c.showPage()
            y = h - 72
    
    # Add sources section
    y -= 12
    if y < 120:
        c.showPage()
        y = h - 72
    c.setFont("Times-Bold", 12)
    c.drawString(72, y, "Sources")
    y -= 24
    c.setFont("Times-Roman", 11)
    for s in sources or []:
        c.drawString(72, y, f"{s.get('book')} | chunk={s.get('chunk_id')} | g={s.get('granularity')} | pos={s.get('position')}")
        y -= 16
        if y < 72:
            c.showPage()
            y = h - 72
    
    c.save()

def export_qa_docx(question, answer, sources, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    d = Document()
    d.add_heading("Question", level=1)
    d.add_paragraph(question)
    d.add_heading("Answer", level=1)
    for line in (answer or "").split("\n"):
        d.add_paragraph(line)
    d.add_heading("Sources", level=2)
    for s in sources or []:
        d.add_paragraph(f"{s.get('book')} | chunk={s.get('chunk_id')} | g={s.get('granularity')} | pos={s.get('position')}")
    d.save(out_path)

def export_quiz_docx(topic, items, sources, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    d = Document()
    if topic:
        d.add_heading(f"Quiz on {topic}", level=1)
    else:
        d.add_heading("Quiz", level=1)

    for i, it in enumerate(items, 1):
        d.add_paragraph(f"{i}. {it.get('question')}")
        opts = it.get("options") or {}
        for k in ["A","B","C","D"]:
            d.add_paragraph(f"{k}. {opts.get(k,'')}", style="List Bullet")
    d.add_page_break()
    d.add_heading("Answer Key", level=1)
    for i, it in enumerate(items, 1):
        d.add_paragraph(f"{i}. {it.get('correct')}")
    d.add_heading("Sources", level=2)
    for s in sources or []:
        d.add_paragraph(f"{s.get('book')} | chunk={s.get('chunk_id')} | g={s.get('granularity')} | pos={s.get('position')}")
    d.save(out_path)

def export_quiz_pptx(topic, items, sources, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs = Presentation()

    if topic: 
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = f"Quiz on {topic}"
        # if slide.placeholders.len > 1:
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = ""

    for i, it in enumerate(items, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Q{i}: {it.get('question')}"
        body = slide.shapes.placeholders[1].text_frame
        body.clear()
        opts = it.get("options") or {}
        for k in ["A","B","C","D"]:
            body.add_paragraph().text = f"{k}. {opts.get(k,'')}"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Answer Key"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, it in enumerate(items, 1):
        body.add_paragraph().text = f"{i}. {it.get('correct')}"
    
    # Add sources slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sources"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for s in sources or []:
        body.add_paragraph().text = f"{s.get('book')} | chunk={s.get('chunk_id')} | g={s.get('granularity')} | pos={s.get('position')}"
    
    prs.save(out_path)

class EmailPayload(BaseModel):
    recipient: str = Field(description="The email address of the recipient.")
    subject: str = Field(description="The subject line for the email.")
    body: str = Field(description="The body content of the email.")
    file_path: Optional[str] = Field(description="Optional path to a file to be sent as an attachment.")

class EmailTool(BaseTool):
    """
    A tool to send an email. 
    Per project spec, this is a placeholder and will log instead of sending.
    """
    name: str = "send_email"
    description: str = "Sends an email with an optional attachment. Use this to send reports or quizzes to users."
    args_schema: Type[BaseModel] = EmailPayload

    def _run(self, recipient: str, subject: str, body: str, file_path: str = None) -> str:
        # In a real system, this would integrate with SMTP or an API (e.g., SendGrid).
        # For this project, we just log the action.
        log_msg = f"Email Tool: \n  To: {recipient}\n  Subject: {subject}\n  Body: {body[:50]}...\n  Attachment: {file_path}"
        logger.info(log_msg)
        print(log_msg)
        
        if file_path and not os.path.exists(file_path):
            return f"Error: Could not send email, attachment not found at {file_path}"
        
        return f"Email successfully queued for {recipient} with subject '{subject}'."

    async def _arun(self, *args, **kwargs):
        # This tool does not support async
        return self._run(*args, **kwargs)


class ExportPayload(BaseModel):
    export_type: str = Field(description="The type of export, e.g., 'qa' or 'quiz'.")
    file_format: str = Field(description="The file format, e.g., 'pdf', 'docx', or 'pptx'.")
    data: Dict[str, Any] = Field(description="The JSON data to be exported. For 'qa', this is {'answer': ..., 'sources': [...]}. For 'quiz', this is {'items': [...]}.")

class ExportTool(BaseTool):
    """
    A tool to generate and export documents (PDF, DOCX, PPTX).
    This fulfills the requirement for Document Generation.
    """
    name: str = "export_document"
    description: str = "Generates and exports a document (PDF, DOCX, PPTX) from QA or Quiz data. Returns the path to the saved file."
    args_schema: Type[BaseModel] = ExportPayload

    def _run(self, export_type: str, file_format: str, data: Dict[str, Any]) -> str:
        exports_dir = "./exports"
        os.makedirs(exports_dir, exist_ok=True)
        timestamp = int(time.time())
        
        try:
            if export_type == "qa":
                question = data.get("question", "")
                answer = data.get("answer", "")
                sources = data.get("sources", [])
                if file_format == "pdf":
                    path = os.path.join(exports_dir, f"qa_export_{timestamp}.pdf")
                    export_qa_pdf(question, answer, sources, path)
                    return path
                elif file_format == "docx":
                    path = os.path.join(exports_dir, f"qa_export_{timestamp}.docx")
                    export_qa_docx(question, answer, sources, path)
                    return path
            
            elif export_type == "quiz":
                topic = data.get("topic", "")
                items = data.get("items", [])
                sources = data.get("sources", [])                
                if not items:
                    return "Error: No quiz items provided in data."
                if file_format == "pdf":
                    path = os.path.join(exports_dir, f"quiz_export_{timestamp}.pdf")
                    export_quiz_pdf(topic, items, sources, path)
                    return path
                elif file_format == "docx":
                    path = os.path.join(exports_dir, f"quiz_export_{timestamp}.docx")
                    export_quiz_docx(topic, items, sources, path)
                    return path
                elif file_format == "pptx":
                    path = os.path.join(exports_dir, f"quiz_export_{timestamp}.pptx")
                    export_quiz_pptx(topic, items, sources, path)
                    return path
                    
            return f"Error: Invalid export_type '{export_type}' or file_format '{file_format}'."
        
        except Exception as e:
            logger.error(f"ExportTool failed: {e}", exc_info=True)
            return f"Error: Failed to export document. {e}"

    async def _arun(self, *args, **kwargs):
        # This tool does not support async
        return self._run(*args, **kwargs)
